"""
Render a ResumeContent into ANY analysed template using its TemplateMap (see analyzer.py).

The renderer walks the original paragraph list in order and, per role:
  * headings / static / blank / drawing  -> cloned (headings get their canonical text)
  * summary                              -> new summary (multi-paragraph summaries collapse into the first)
  * skill_line                           -> the first prototype emits ALL skill lines, later ones are skipped
  * experience zone                      -> the learned job block is emitted once per job
  * edu_line                             -> filled from its pattern
Header-like lines (name, contact, job headers, scope, url, education) are rebuilt from their PATTERN:
placeholders take content values, every token borrows the run formatting of the original text it replaces.
"""
from __future__ import annotations

import copy
import re
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..models import Experience, ResumeContent
from .engine import BR, TAB, Para, TemplateDoc, donor_rpr_by_text, md_segments, rewrite_inplace, strip_md

SEP_RE = re.compile(r"^[\s|·•,;:–—\-/()\[\]]*$")
TOKEN_RE = re.compile(r"\{(\w+)\}")


# --------------------------------------------------------------------------- variables
def _dates(start: str, end: str, original: str | None) -> str:
    sep = " – "
    if original:
        m = re.search(r"\s*(–|—|-|to|until)\s*", original)
        if m:
            sep = m.group(0)
            if not sep.startswith(" "):
                sep = " " + sep.strip() + " " if sep.strip() in ("to", "until") else sep
    return f"{start}{sep}{end}"


def _vars_header(c: ResumeContent) -> dict[str, str]:
    ct = c.contact
    return {"name": c.full_name, "headline": c.headline, "location": c.location, "email": ct.email,
            "phone": ct.phone, "linkedin": ct.linkedin, "github": ct.github, "website": ct.website}


def _vars_exp(e: Experience, values: dict) -> dict[str, str]:
    return {"company": e.company, "location": e.location, "remote": "Remote" if e.remote else "", "role": e.role,
            "start": e.start, "end": e.end, "dates": _dates(e.start, e.end, values.get("dates")),
            "tech_scope": e.tech_scope, "website": e.website}


def _vars_edu(c: ResumeContent, values: dict) -> dict[str, str]:
    ed = c.education
    yrs = ""
    if ed.start_year and ed.end_year:
        yrs = _dates(ed.start_year, ed.end_year, values.get("years"))
    elif ed.end_year or ed.start_year:
        yrs = ed.end_year or ed.start_year
    return {"degree": ed.degree, "field": ed.field, "university": ed.university,
            "start_year": ed.start_year, "end_year": ed.end_year, "years": yrs}


# --------------------------------------------------------------------------- pattern -> segments
def pattern_segments(proto, pattern: str, values: dict, vars: dict[str, str]) -> list:
    """Tokenise the pattern; drop empty placeholders together with one adjacent separator."""
    pattern = pattern.replace("\\t", "\t").replace("\\n", "\n")
    tokens: list[dict] = []
    pos = 0
    for m in TOKEN_RE.finditer(pattern):
        if m.start() > pos:
            tokens.append({"lit": pattern[pos:m.start()]})
        tokens.append({"key": m.group(1)})
        pos = m.end()
    if pos < len(pattern):
        tokens.append({"lit": pattern[pos:]})
    if not tokens:
        return [(pattern, "")]

    # resolve values, mark empties
    for t in tokens:
        if "key" in t:
            t["text"] = vars.get(t["key"], "")
            t["orig"] = str(values.get(t["key"], "") or "")
            # follow the template's casing convention (e.g. "ALEX MORGAN" -> uppercase names)
            letters = re.sub(r"[^A-Za-z]", "", t["orig"])
            if len(letters) >= 3 and letters.isupper() and t["key"] in ("name", "company", "university", "role"):
                t["text"] = t["text"].upper()
    # remove empty placeholders + one separator neighbour
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if "key" in t and not t["text"]:
            nxt = tokens[i + 1] if i + 1 < len(tokens) else None
            prv = tokens[i - 1] if i > 0 else None
            del tokens[i]
            if nxt is not None and "lit" in nxt and SEP_RE.match(nxt["lit"]) and nxt["lit"].strip():
                tokens.remove(nxt)
            elif prv is not None and "lit" in prv and SEP_RE.match(prv["lit"]) and prv["lit"].strip():
                tokens.remove(prv)
                i -= 1
            continue
        i += 1
    # collapse doubled separators left behind (e.g. " | " " | ")
    j = 1
    while j < len(tokens):
        a, b = tokens[j - 1], tokens[j]
        if "lit" in a and "lit" in b:
            a["lit"] = a["lit"] + b["lit"]
            del tokens[j]
            continue
        j += 1
    for t in tokens:
        if "lit" in t and SEP_RE.match(t["lit"]):
            t["lit"] = re.sub(r"(\s*[|·•]\s*){2,}", lambda m: m.group(1), t["lit"])
    # trailing/leading pure separators
    while tokens and "lit" in tokens[0] and SEP_RE.match(tokens[0]["lit"]) and tokens[0]["lit"].strip(" ") and "\t" not in tokens[0]["lit"]:
        if tokens[0]["lit"].strip() in ("(", "["):
            break
        tokens.pop(0)
    while tokens and "lit" in tokens[-1] and SEP_RE.match(tokens[-1]["lit"]) and tokens[-1]["lit"].strip() and tokens[-1]["lit"].strip() not in (")", "]", "|"):
        tokens.pop()

    segs: list = []
    last_rpr = None
    for t in tokens:
        if "key" in t:
            rpr = donor_rpr_by_text(proto, t["orig"]) if t["orig"] else _main_run_rpr(proto)
            segs.append((t["text"], rpr))
            last_rpr = rpr
        else:
            lit = t["lit"]
            if lit.strip():
                rpr = donor_rpr_by_text(proto, lit.strip())
            else:
                rpr = last_rpr if last_rpr is not None else donor_rpr_by_text(proto, "")
            segs.append((lit, rpr))
            last_rpr = rpr
    return segs


def _main_run_rpr(proto):
    """rPr of the longest text run that is not a pure separator (the paragraph's 'main' content)."""
    import copy as _copy
    best, best_len = None, -1
    for r in proto.iter(qn("w:r")):
        txt = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if txt.strip() and not SEP_RE.match(txt) and len(txt.strip()) > best_len:
            best, best_len = r, len(txt.strip())
    if best is None:
        return donor_rpr_by_text(proto, "")
    rpr = best.find(qn("w:rPr"))
    return _copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")


def _add_right_tab(p, pos_twips: int):
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    tabs = ppr.get_or_add_tabs()
    for t in list(tabs):
        tabs.remove(t)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(pos_twips))
    tabs.append(tab)


def _text_width_twips(page: dict) -> int:
    w = page.get("width_in") or 8.5
    m = page.get("margins_in") or [1, 1, 1, 1]
    left, right = (m[2] or 1), (m[3] or 1)
    return int((w - left - right) * 1440) - 20


# --------------------------------------------------------------------------- renderer
class GenericTemplate:
    def __init__(self, tmap: dict, docx_path: str):
        self.map = tmap
        self.path = docx_path

    def _prepare(self, td: TemplateDoc):
        """Register every body paragraph as a prototype by index and fix tab alignment."""
        for p in self.map["paragraphs"]:
            td.register(f"p{p['idx']}", p["idx"])
            # header lines that used a run of tabs to push the date right -> one right tab stop
            if p["role"] in ("exp_header_1", "exp_header_2", "exp_header_3", "name_line", "contact_line", "edu_line"):
                proto = td.protos[f"p{p['idx']}"]
                ntabs = sum(1 for r in proto.iter(qn("w:tab")))
                if ntabs >= 2 and "\t" in (p.get("pattern") or "").replace("\\t", "\t"):
                    _add_right_tab(proto, _text_width_twips(self.map["page"]))
                    p["_single_tab"] = True

    def _pat(self, p: dict, vars: dict, td: TemplateDoc) -> Para:
        proto = td.protos[f"p{p['idx']}"]
        pattern = p.get("pattern") or self.map["texts"][str(p["idx"])]
        if p.get("_single_tab"):
            pattern = re.sub(r"[ \t]*(?:\\t|\t)[ \t\\]*", "\t", pattern.replace("\\t", "\t"))
        segs = pattern_segments(proto, pattern, p.get("values") or {}, vars)
        return Para(f"p{p['idx']}", segs)

    def compose(self, td: TemplateDoc, c: ResumeContent):
        self._prepare(td)
        paras = self.map["paragraphs"]
        exp = self.map["experience"]
        hv = _vars_header(c)
        keep_first = 1 if paras and paras[0]["role"] == "drawing" else 0
        td.clear_body(keep_first_n=keep_first)
        self._fill_textboxes(td, c)

        A = td.append
        i = 0
        skills_done = False
        summary_done = False
        n = len(paras)
        while i < n:
            p = paras[i]
            role = p["role"]
            key = f"p{p['idx']}"
            if role == "drawing":
                if i != 0:
                    td.body.find(qn("w:sectPr")).addprevious(copy.deepcopy(td.protos[key]))
                i += 1
                continue
            if role in ("name_line", "contact_line"):
                A(self._pat(p, hv, td))
            elif role.startswith("heading_"):
                orig = self.map["texts"][str(p["idx"])].strip()
                text = p.get("heading_text") or orig
                if orig and re.sub(r"[^A-Za-z]", "", orig).isupper():
                    text = text.upper()
                A(Para.plain(key, text))
            elif role == "summary":
                if not summary_done:
                    lead = re.match(r"^\s*", self.map["texts"][str(p["idx"])]).group(0)
                    A(Para.plain(key, lead + strip_md(c.summary)))
                    summary_done = True
            elif role == "skill_line":
                if not skills_done:
                    labelled = self.map["budget"]["skill_labelled"]
                    for s in c.skills:
                        if labelled:
                            label = s.label or "Skills"
                            A(Para(key, [(label, "b"), (": " + s.items, "")]))
                        else:
                            A(Para.plain(key, s.items))
                    skills_done = True
            elif role == "edu_line":
                ev = _vars_edu(c, p.get("values") or {})
                pat = p.get("pattern") or ""
                if "{degree}" in pat and "{field}" not in pat and c.education.field:
                    # pattern has no separate field slot: fold it into the degree ("Bachelor of Computer Science")
                    deg = c.education.degree.strip()
                    ev["degree"] = (f"{deg} of {c.education.field}" if deg.lower() in ("bachelor", "master", "bachelor's", "master's", "bachelor’s", "master’s")
                                    else f"{deg}, {c.education.field}" if " of " not in deg.lower() else f"{deg} in {c.education.field}")
                A(self._pat(p, ev, td))
            elif role in ("blank", "static"):
                if role == "blank":
                    A(Para(key))
                else:
                    A(Para.plain(key, self.map["texts"][str(p["idx"])]))
            elif role.startswith("exp_"):
                if i == exp["first"]:
                    self._emit_jobs(td, c)
                    i = exp["end"]
                    continue
                # stray exp_* paragraph outside the learned zone: skip
            i += 1

    def _emit_jobs(self, td: TemplateDoc, c: ResumeContent):
        block = self.map["experience"]["block"]
        by_role: dict[str, list[dict]] = {}
        for b in block:
            by_role.setdefault(b["role"], []).append(b)
        A = td.append
        for e in c.experiences:
            bullets_emitted = False
            bullet_i = 0
            for b in block:
                role = b["role"]
                key = f"p{b['idx']}"
                pmap = self.map["paragraphs"][b["idx"]]
                ev = _vars_exp(e, pmap.get("values") or {})
                if role in ("exp_header_1", "exp_header_2", "exp_header_3"):
                    A(self._pat(pmap, ev, td))
                elif role == "exp_scope":
                    if e.tech_scope:
                        A(self._pat(pmap, ev, td))
                elif role == "exp_blurb":
                    if e.company_blurb:
                        A(Para.plain(key, strip_md(e.company_blurb)))
                elif role == "exp_subheading":
                    A(Para.plain(key, self.map["texts"][str(b["idx"])].strip()))
                elif role == "exp_bullet":
                    if not bullets_emitted:
                        for bl in e.bullets:
                            if self.map["budget"]["bold_keywords"]:
                                A(Para(key, md_segments(bl)))
                            else:
                                A(Para.plain(key, strip_md(bl)))
                        bullets_emitted = True
                    bullet_i += 1
                elif role == "exp_url":
                    if e.website:
                        A(self._pat(pmap, ev, td))
                elif role == "blank":
                    A(Para(key))
                elif role == "static":
                    A(Para.plain(key, self.map["texts"][str(b["idx"])]))

    def _fill_textboxes(self, td: TemplateDoc, c: ResumeContent):
        tb = self.map.get("textbox_paragraphs") or []
        if not any(p["role"] in ("name_line", "contact_line") for p in tb):
            return
        hv = _vars_header(c)
        seen: dict[str, list] = {}
        # every copy (DrawingML + VML fallback) of every text box, in document order
        for txbx in td.body.iter(qn("w:txbxContent")):
            key = "".join(t.text or "" for t in txbx.iter(qn("w:t")))
            paras = txbx.findall(qn("w:p"))
            seen.setdefault(key, []).append(paras)
        # map order = first-seen order of distinct boxes
        offset = 0
        for key, copies in seen.items():
            count = len(copies[0])
            for j in range(count):
                idx = offset + j
                if idx >= len(tb):
                    break
                p = tb[idx]
                if p["role"] not in ("name_line", "contact_line"):
                    continue
                for paras in copies:
                    if j < len(paras):
                        proto = paras[j]
                        pattern = p.get("pattern") or self.map["textbox_texts"][str(idx)]
                        segs = pattern_segments(proto, pattern, p.get("values") or {}, hv)
                        rewrite_inplace(proto, segs)
            offset += count


def render_generic(tmap: dict, docx_path: str, content: ResumeContent, out_path: str) -> str:
    td = TemplateDoc(docx_path)
    GenericTemplate(copy.deepcopy(tmap), docx_path).compose(td, content)
    td.save(out_path)
    return out_path
